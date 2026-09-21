"""Convert REPORT.md to REPORT.docx (headings, paragraphs, lists, pipe tables, quotes, code blocks, images). python-docx only.
Usage: .venv/bin/python md_to_docx.py [in.md] [out.docx]"""
import re, sys
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`)")

def shade(el, fill):
    pr = el.get_or_add_pPr() if hasattr(el, "get_or_add_pPr") else el
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill); pr.append(shd)

def border(p, side, color="2F5496", sz="18"):
    pPr = p._p.get_or_add_pPr(); b = OxmlElement("w:pBdr"); e = OxmlElement(f"w:{side}")
    e.set(qn("w:val"), "single"); e.set(qn("w:sz"), sz); e.set(qn("w:space"), "4"); e.set(qn("w:color"), color); b.append(e); pPr.append(b)

def runs(p, text, size=None, bold=False, italic=False):
    for part in INLINE.split(text):
        if not part:
            continue
        b, i, code = bold, italic, False
        if part.startswith("**") and part.endswith("**"): part, b = part[2:-2], True
        elif part.startswith("`") and part.endswith("`"): part, code = part[1:-1], True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2: part, i = part[1:-1], True
        r = p.add_run(part); r.bold, r.italic = b, i
        if code: r.font.name, r.font.size = "Consolas", Pt(9.5)
        elif size: r.font.size = Pt(size)

def set_cell_shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill); tcPr.append(shd)

def add_table(doc, rows):
    header, body = rows[0], rows[2:]
    n = len(header); t = doc.add_table(rows=1 + len(body), cols=n); t.style = "Table Grid"; t.autofit = False
    lens = [max(len(re.sub(r"[*`]", "", r[c])) if c < len(r) else 0 for r in [header] + body) for c in range(n)]
    w = [max(min(l, 60), 8) for l in lens]; tot = sum(w)
    for ri, row in enumerate([header] + body):
        for ci in range(n):
            cell = t.cell(ri, ci); cell.width = Cm(16.2 * w[ci] / tot); cell.text = ""
            p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(1)
            runs(p, row[ci].strip() if ci < len(row) else "", size=9, bold=(ri == 0))
            if ri == 0: set_cell_shade(cell, "DBE5F1")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def convert(src, dst):
    doc = Document(); s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7); s.left_margin = s.right_margin = Cm(2.4); s.top_margin = s.bottom_margin = Cm(2.2)
    st = doc.styles["Normal"]; st.font.name, st.font.size = "Calibri", Pt(11); st.paragraph_format.space_after = Pt(6)
    for name, size in (("Heading 1", 22), ("Heading 2", 16), ("Heading 3", 13)):
        h = doc.styles[name]; h.font.name, h.font.size, h.font.color.rgb = "Calibri", Pt(size), RGBColor(0x1F, 0x38, 0x64)
    lines = open(src, encoding="utf-8").read().split("\n"); i = 0
    while i < len(lines):
        ln = lines[i]
        if not ln.strip(): i += 1; continue
        if ln.startswith("```"):
            code = []; i += 1
            while not lines[i].startswith("```"): code.append(lines[i]); i += 1
            i += 1; p = doc.add_paragraph(); shade(p._p, "F3F3F3"); r = p.add_run("\n".join(code)); r.font.name, r.font.size = "Consolas", Pt(9); continue
        m = re.match(r"(#{1,3}) (.*)", ln)
        if m: doc.add_heading(re.sub(r"[*`]", "", m.group(2)), level=len(m.group(1))); i += 1; continue
        if ln.strip() == "---": p = doc.add_paragraph(); border(p, "bottom", "999999", "6"); i += 1; continue
        m = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", ln)
        if m:
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(m.group(2), width=Inches(6.0)); i += 1; continue
        if ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"): rows.append([c for c in lines[i].strip().strip("|").split("|")]); i += 1
            add_table(doc, rows); continue
        if ln.startswith(">"):
            while i < len(lines) and lines[i].startswith(">"):
                txt = lines[i][1:].strip(); i += 1
                if not txt: continue
                p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.4); p.paragraph_format.space_after = Pt(2); shade(p._p, "EEF3FA"); border(p, "left")
                runs(p, txt)
            doc.paragraphs[-1].paragraph_format.space_after = Pt(8); continue
        m = re.match(r"(\s*)(- |\d+\. )(.*)", ln)
        if m:
            while i < len(lines) and (m := re.match(r"(\s*)(- |\d+\. )(.*)", lines[i])):
                p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.9); p.paragraph_format.first_line_indent = Cm(-0.5); p.paragraph_format.space_after = Pt(3)
                runs(p, ("•  " if m.group(2) == "- " else m.group(2) + " ") + m.group(3)); i += 1
            continue
        para = []
        while i < len(lines) and lines[i].strip() and not re.match(r"(#{1,3} |```|\||>|!\[|---$|\s*(- |\d+\. ))", lines[i]): para.append(lines[i]); i += 1
        p = doc.add_paragraph()
        for k, t in enumerate(para):
            if k: p.add_run().add_break()
            runs(p, t)
        if not para: i += 1
    doc.save(dst)

if __name__ == "__main__":
    convert(*(sys.argv[1:3] or ["REPORT.md", "REPORT.docx"])); print("wrote", sys.argv[2] if len(sys.argv) > 2 else "REPORT.docx")
